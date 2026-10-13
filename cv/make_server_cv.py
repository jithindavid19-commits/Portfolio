import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1]
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.2)
    s.left_margin = s.right_margin = Cm(1.6)
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)  # A4

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.2

def runs(p, text, bold=False, size=None):
    # [[...]] = highlight in yellow (details to fill in / confirm)
    for i, part in enumerate(re.split(r"\[\[(.*?)\]\]", text)):
        if not part: continue
        r = p.add_run(part); r.bold = bold
        if size: r.font.size = Pt(size)
        if i % 2: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def link(p, text, url):
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = p.add_run(text); h.append(r._r); p._p.append(h)

def para(text="", bold=False, size=None, after=0, before=0):
    p = doc.add_paragraph(); runs(p, text, bold, size)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    return p

def rule(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "2", "w:color": "333333"}.items(): bt.set(qn(k), v)
    b.append(bt); pPr.append(b)

def heading(t):
    p = para(t.upper(), bold=True, size=11.5, before=15, after=6); rule(p)

def job(title, place, dates, bullets):
    p = para(before=10, after=3)
    runs(p, title, bold=True); runs(p, " – " + place)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), alignment=2)
    runs(p, "\t" + dates)
    for b in bullets:
        q = doc.add_paragraph(style="List Bullet"); runs(q, b)
        q.paragraph_format.space_after = Pt(1.5)

name = para("JITHIN GEORGE", bold=True, size=20, after=3)
para("Wembley, London HA9 0TT")
p = para("+44 7721 960626  |  ")
link(p, "jithindavid.19@gmail.com", "mailto:jithindavid.19@gmail.com")
runs(p, "  |  ")
link(p, "linkedin.com/in/jithin-george-jj1999", "https://www.linkedin.com/in/jithin-george-jj1999")
para("Available for Full-Time and Part-Time Roles", after=2)

heading("Professional Summary")
para("Front of house professional with 2.5 years in busy casual dining, street food and "
     "restaurant settings, most recently as a Restaurant Supervisor. Takes full ownership of a section from set-up to "
     "close, reads guests' needs early and gives every table an attentive, informed and personalised experience. Calm "
     "under pressure, quick with the kitchen and bar, and always ready to step in for the team at peak times.")

heading("Employment History")
job("Restaurant Supervisor", "Tru Street Food, UK", "Mar 2025 – Sep 2026", [
    "Ran front of house across busy lunch and dinner services, keeping service fast, consistent and friendly.",
    "Led pre-shift briefings on specials, allergens and menu changes, and allocated sections to the team.",
    "Trained and supported new and junior team members on service standards, the menu and the EPOS.",
    "Handled guest feedback and complaints professionally, escalating to management when appropriate.",
    "Kept kitchen and front of house in sync during peak hours so food went out hot, accurate and on time.",
    "Oversaw opening and closing, cashing up, stock checks and food safety compliance.",
])
job("Front of House Staff", "Bardez, UK", "Aug 2024 – Mar 2025", [
    "Greeted and seated guests, managed reservations and walk-ins, and kept table turnover smooth.",
    "Owned a section of 6–8 tables, pacing each course with the kitchen for a relaxed guest experience.",
    "Explained dishes, ingredients and spice levels, and recommended wine and cocktail pairings.",
    "Upsold specials, sides and desserts to increase average spend per table.",
    "Handled allergen and dietary requests (vegan, gluten-free, nut allergies) accurately with the kitchen.",
    "Took orders and payments accurately on the EPOS, including split bills, card and cash.",
    "Completed side duties, table settings, polishing and restocking so the section was ready for every shift.",
    "Built rapport with regular guests and resolved feedback calmly, encouraging repeat bookings.",
])
job("Front of House Staff", "Chit n Chaat, UK", "Feb 2024 – Jul 2024", [
    "Served eat-in and takeaway guests during high-volume shifts with speed and warmth.",
    "Guided guests through the menu and handled special requests, helping them order with confidence.",
    "Stepped in to support teammates in rush periods, cutting wait times and keeping service flowing.",
    "Passed orders clearly to the kitchen so food went out quickly and accurately at busy times.",
    "Kept the section clean, restocked and well presented in line with food safety standards.",
])

heading("Education")
p = para(after=1); runs(p, "MSc Digital Marketing", bold=True); runs(p, " – University of Salford  |  2024 – 2025")
para("Dissertation: The Impact of Modern Influencer Strategies on Consumer Behaviour – Distinction (71.33%)", after=4)
p = para(); runs(p, "BA Mass Media and Communication", bold=True); runs(p, " – University of Mumbai, India  |  2018 – 2021")

heading("Key Skills")
skills = ["Section management, set-up to close", "Menu knowledge and recommendations",
          "Allergen and dietary awareness", "Upselling food, wine and cocktails",
          "Orders, payments and EPOS", "Complaint handling and escalation",
          "Licensing, health and safety, food hygiene", "Training and supporting junior staff"]
t = doc.add_table(rows=4, cols=2)
for i, sk in enumerate(skills):
    c = t.cell(i % 4, i // 4); c.paragraphs[0].style = doc.styles["List Bullet"]; runs(c.paragraphs[0], sk)

doc.save(out)

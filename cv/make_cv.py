import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1]
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.3)
    s.left_margin = s.right_margin = Cm(1.9)
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)  # A4

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.05

def runs(p, text, bold=False, size=None):
    # [[...]] = highlight in yellow (details to fill in / confirm)
    for i, part in enumerate(re.split(r"\[\[(.*?)\]\]", text)):
        if not part: continue
        r = p.add_run(part); r.bold = bold
        if size: r.font.size = Pt(size)
        if i % 2: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def para(text="", bold=False, size=None, after=0, before=0):
    p = doc.add_paragraph(); runs(p, text, bold, size)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    return p

def rule(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "2", "w:color": "333333"}.items(): bt.set(qn(k), v)
    b.append(bt); pPr.append(b)

def heading(t):
    p = para(t.upper(), bold=True, size=11.5, before=9, after=4); rule(p)

def job(title, place, dates, bullets):
    p = para(before=7, after=2)
    runs(p, title, bold=True); runs(p, " – " + place)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), alignment=2)
    runs(p, "\t" + dates)
    for b in bullets:
        q = doc.add_paragraph(style="List Bullet"); runs(q, b)
        q.paragraph_format.space_after = Pt(1)

name = para("JITHIN GEORGE", bold=True, size=20, after=3)
para("Wembley Park, London HA9 0TT")
para("+44 7721 960626  |  jithindavid.19@gmail.com  |  linkedin.com/in/jithin-george-jj1999")
para("Available for Full-Time and Part-Time Roles  |  Flexible across Weekends, Evenings and Bank Holidays", after=2)

heading("Professional Summary")
para("Friendly, dependable retail and customer service professional with UK experience on the front of house "
     "and in a busy warehouse, backed by three years in client-facing marketing roles. Comfortable on the shop "
     "floor, at the till and in the stockroom, with a good eye for product presentation and a habit of keeping "
     "track of the numbers that matter – sales, stock and customer feedback. MSc Digital Marketing graduate "
     "(dissertation on consumer behaviour) who understands what makes customers buy and come back.")

heading("Retail & Customer Service Experience")
job("Front of House Team Member (Part-time)", "[[Venue name]], Manchester", "[[Month 2024 – Month 2025]]", [
    "Served [[150+]] customers per shift at peak times, keeping queues moving while staying warm and welcoming.",
    "Took cash and card payments on the EPOS till with [[zero]] discrepancies at end-of-shift cash-ups.",
    "Suggested add-ons and current promotions to help increase average spend per customer.",
    "Dealt with complaints and queries calmly on the spot, passing to a manager only when needed.",
    "Kept the customer area clean, well stocked and in line with health, safety and food hygiene rules.",
])
job("Warehouse Operative (Part-time)", "[[Company name]], Manchester", "[[Month 2024 – Month 2025]]", [
    "Picked, packed and dispatched [[100+]] orders per shift with a handheld scanner, meeting daily pick targets.",
    "Checked incoming deliveries against delivery notes and reported damaged or missing items.",
    "Supported weekly stock counts and replenishment so products were accurate and easy to locate.",
    "Followed manual handling and health and safety procedures in a fast-paced environment.",
])

heading("Other Experience")
job("Marketing Executive", "Qyuki (Remote)", "Mar 2024 – Present", [
    "Look after 8–12 client accounts each campaign cycle, meeting 100% of delivery targets through planning and regular follow-ups.",
    "Grew the partner contact list from 300 to 450+, opening up 25% more partnership opportunities.",
    "Ran 4 campaigns from first brief to final delivery, keeping clients informed at every stage.",
    "Track KPIs in Excel and Trello and create client presentations in Canva.",
])
job("Influencer Marketing Executive", "Tring", "Jul 2022 – Apr 2023", [
    "Onboarded 10+ creators per campaign, managing contracts, timelines and deliverables.",
    "Built and maintained a database of 500+ creators, making it quicker to match brands with the right people.",
    "Reviewed results across 60+ posts (reach, impressions, engagement) to choose top performers for repeat work.",
])
job("Talent Coordinator", "Aspiring Productions (Reality Show)", "May 2022 – Dec 2022", [
    "Sourced 300+ applicants through social media and ran online auditions for 250+ participants.",
    "Worked with a 4–6 person production team to organise audition schedules and availability.",
    "Set up a shared Excel tracker that cut turnaround time by 40%.",
])

heading("Education")
p = para(after=1); runs(p, "MSc Digital Marketing", bold=True); runs(p, " – University of Salford, Manchester  |  2025")
para("Dissertation: The Impact of Modern Influencer Strategies on Consumer Behaviour – Distinction (71.33%)", after=4)
p = para(); runs(p, "BA Mass Media and Communication", bold=True); runs(p, " – University of Mumbai, India  |  2021")

heading("Key Skills")
skills = ["Customer service and complaint handling", "EPOS, cash handling and card payments",
          "Sales, upselling and product knowledge", "Visual merchandising and shop-floor standards",
          "Stock control, deliveries and replenishment", "Handheld scanners and stock counts",
          "Microsoft Excel, Google Sheets and Canva", "Teamwork and working under pressure"]
t = doc.add_table(rows=4, cols=2)
for i, s in enumerate(skills):
    c = t.cell(i % 4, i // 4); c.paragraphs[0].text = ""
    c.paragraphs[0].style = doc.styles["List Bullet"]; runs(c.paragraphs[0], s)

para("References available on request.", before=8)
doc.save(out)

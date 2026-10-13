import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1]
doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(1.2)
    s.left_margin = s.right_margin = Cm(1.6)
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)  # A4

st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.2

def runs(p, text, bold=False, size=None):
    # [[...]] = highlight in yellow (details to fill in / confirm)
    for i, part in enumerate(re.split(r"\[\[(.*?)\]\]", text)):
        if not part: continue
        r = p.add_run(part); r.bold = bold
        if size: r.font.size = Pt(size)
        if i % 2: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def link(p, text, url):
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = p.add_run(text); h.append(r._r); p._p.append(h)

def para(text="", bold=False, size=None, after=0, before=0):
    p = doc.add_paragraph(); runs(p, text, bold, size)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    return p

def rule(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "2", "w:color": "333333"}.items(): bt.set(qn(k), v)
    b.append(bt); pPr.append(b)

def heading(t):
    p = para(t.upper(), bold=True, size=11.5, before=15, after=6); rule(p)

def job(title, place, dates, bullets):
    p = para(before=10, after=3)
    runs(p, title, bold=True); runs(p, " – " + place)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.8), alignment=2)
    runs(p, "\t" + dates)
    for b in bullets:
        q = doc.add_paragraph(style="List Bullet"); runs(q, b)
        q.paragraph_format.space_after = Pt(1.5)

name = para("JITHIN GEORGE", bold=True, size=20, after=3)
para("Wembley, London HA9 0TT")
p = para("+44 7721 960626  |  ")
link(p, "jithindavid.19@gmail.com", "mailto:jithindavid.19@gmail.com")
runs(p, "  |  ")
link(p, "linkedin.com/in/jithin-george-jj1999", "https://www.linkedin.com/in/jithin-george-jj1999")
para("Available for Full-Time and Part-Time Roles", after=2)

heading("Professional Summary")
para("Front of house professional with 2.5 years in fast-paced street food, casual dining and restaurant settings, "
     "most recently as a Restaurant Supervisor. Friendly, upbeat and quick on the counter, with a real passion for fresh, "
     "high-quality food and making sure every guest leaves with a smile. Keen to keep learning across front of house "
     "and kitchen, and to grow towards Shift Leader in a busy, team-focused brand.")

heading("Employment History")
job("Restaurant Supervisor", "Tru Street Food, UK", "Mar 2025 – Sep 2026", [
    "Ran front of house across busy lunch and dinner services, keeping service fast, friendly and consistent.",
    "Kept kitchen and counter in sync during peak hours so fresh food went out hot, accurate and on time.",
    "Trained and supported new team members on service standards, the menu, allergens and the EPOS.",
    "Led pre-shift briefings and allocated roles so the whole team knew their part before every rush.",
    "Handled guest feedback and complaints professionally, escalating to management when appropriate.",
    "Oversaw opening and closing, cashing up, stock checks and food safety compliance.",
])
job("Shift Supervisor", "Bardez, UK", "Aug 2024 – Mar 2025", [
    "Ran the floor during shifts, allocating sections and keeping service on track at peak times.",
    "Coached new team members on service standards, allergens and side duties during their first shifts.",
    "Welcomed guests warmly, managed reservations and walk-ins, and kept table turnover smooth.",
    "Explained dishes, ingredients and spice levels, helping guests choose with confidence.",
    "Handled allergen and dietary requests (vegan, gluten-free, nut allergies) accurately with the kitchen.",
    "Took orders and payments accurately on the EPOS, including split bills, card and cash.",
    "Completed side duties, cleaning and restocking so the restaurant was ready for every shift.",
])
job("Front of House Staff", "Chit n Chaat, UK", "Feb 2024 – Jul 2024", [
    "Served eat-in and takeaway guests during high-volume shifts with speed and a smile.",
    "Guided guests through the menu and handled special requests, helping them order with confidence.",
    "Passed orders clearly to the kitchen so food went out quickly and accurately at busy times.",
    "Stepped in to support teammates in rush periods, cutting wait times and keeping service flowing.",
    "Kept the counter clean, restocked and well presented in line with food safety standards.",
])

heading("Education")
p = para(after=1); runs(p, "MSc Digital Marketing", bold=True); runs(p, " – University of Salford  |  2024 – 2025")
para("Dissertation: The Impact of Modern Influencer Strategies on Consumer Behaviour – Distinction (71.33%)", after=4)
p = para(); runs(p, "BA Mass Media and Communication", bold=True); runs(p, " – University of Mumbai, India  |  2018 – 2021")

heading("Key Skills")
skills = ["Fast, friendly counter service", "Fresh food and food safety standards",
          "Allergen and dietary awareness", "Orders, payments and EPOS",
          "Teamwork during peak hours", "Training new team members",
          "Opening, closing and cashing up", "Keen to learn and cross-train"]
t = doc.add_table(rows=4, cols=2)
for i, sk in enumerate(skills):
    c = t.cell(i % 4, i // 4); c.paragraphs[0].style = doc.styles["List Bullet"]; runs(c.paragraphs[0], sk)

doc.save(out)
